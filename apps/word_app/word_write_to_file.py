import os
import fire
from docx import Document
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
from action_args import build_command

DEMO = (
    "write text to a word file: "
    "{'app': 'word', 'action': 'write_to_file', 'file_path': [THE_PATH_TO_THE_WORD_FILE], 'contents': [THE_CONTENTS_YOU_WISH_TO_WRITE]}"
)


def construct_action(work_dir, args: dict, py_file_path='/apps/word_app/word_write_to_file.py'):
    return build_command(py_file_path, file_path=args['file_path'], contents=args['contents'])


def word_write_to_file(file_path, contents: str, style='pure-text'):
    try:
        if os.path.exists(file_path):
            document = Document(file_path)
        else:
            document = Document()
        if style == 'pure-text':
            document.add_paragraph(contents)
        elif style == 'title':
            document.add_heading(contents, 0)
        elif style == 'subtitle':
            document.add_heading(contents, 1)
        else:
            raise NotImplementedError
        document.save(file_path)
        return True
    except Exception:
        return False


def main(file_path, contents, style='pure-text'):
    if not os.path.exists(file_path):
        return f'OBSERVATION: The file {file_path} does not exist. Failed to write to the file.'
    success = word_write_to_file(file_path, contents, style)
    if success:
        return 'OBSERVATION: ' + f'Successfully write contents to {file_path}'
    return 'OBSERVATION: ' + f'Failed to write contents to {file_path}'


if __name__ == '__main__':
    fire.Fire(main)
