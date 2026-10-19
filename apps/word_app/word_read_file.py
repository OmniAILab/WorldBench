import os
import fire
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
from action_args import build_command

DEMO = (
    "read the content of the word file: "
    "{'app': 'word', 'action': 'read_file', 'file_path': [THE_PATH_TO_THE_WORD_FILE]}"
)


def construct_action(work_dir, args: dict, py_file_path='/apps/word_app/word_read_file.py'):
    return build_command(py_file_path, file_path=args["file_path"])

def _table_lines(table):
    lines = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        if any(cells):
            lines.append(' | '.join(cells))
    return lines


def _document_lines(doc):
    lines = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith('}p'):
            lines.append(Paragraph(child, doc).text)
        elif child.tag.endswith('}tbl'):
            lines.extend(_table_lines(Table(child, doc)))
    for section in doc.sections:
        for part in (section.header, section.footer):
            lines.extend(p.text for p in part.paragraphs)
            for table in part.tables:
                lines.extend(_table_lines(table))
    return [line for line in lines if line is not None]


def word_read_file(file_path):
    return _document_lines(Document(file_path))

def word_read_file_into_string(file_path):
    return '\n'.join(_document_lines(Document(file_path)))

def wrap(lines):
    ob = "The following is the content from the word file:"
    for line in lines:
        ob += f"\n{line}"
    return ob

def main(file_path, debug=False):
    if not os.path.exists(file_path):
        return f"OBSERVATION: The file {file_path} does not exist. Failed to read the file."
    lines = word_read_file(file_path)
    if debug:
        print(lines)
    return 'OBSERVATION: ' + wrap(lines)


if __name__ == '__main__':
    fire.Fire(main)
