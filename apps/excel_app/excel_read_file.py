import os
import fire
import openpyxl
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
from action_args import build_command

DEMO = (
    "read the excel file to see the existing contents: "
    "{'app': 'excel', 'action': 'read_file', 'file_path': [THE_PATH_TO_THE_EXCEL_FILE]}"
)

def construct_action(work_dir, args: dict, py_file_path='/apps/excel_app/excel_read_file.py'):
    return build_command(py_file_path, file_path=args["file_path"],
                         sheet=args.get("sheet") or args.get("sheet_name"))


def excel_read_file(file_path, sheet=None):
    if sheet is None:
        sheet = openpyxl.load_workbook(file_path).active
    else:
        sheet = openpyxl.load_workbook(file_path)[sheet]
    content_string = ''
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if value is None:
                value = '[Empty Cell]'
            content_string += f'({cell.row}, {cell.column}): {value}\t'
        content_string += '\n'
    return content_string


def wrap(sheet):
    return ("OBSERVATION: The following is the table from the excel file:\n"
            f"{sheet}")

def main(file_path, sheet=None, debug=False):
    if not os.path.exists(file_path):
        return f"OBSERVATION: The file {file_path} does not exist. Failed to read the file."
    contents = excel_read_file(file_path, sheet)
    if debug:
        print(contents)
    return wrap(contents)


if __name__ == '__main__':
    fire.Fire(main)
