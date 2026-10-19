import os
import fire
import openpyxl
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
from action_args import build_command

DEMO = (
    "write text to a cell in the excel file (index starts from 1): "
    "{'app': 'excel', 'action': 'set_cell', 'file_path': [THE_PATH_TO_THE_EXCEL_FILE], 'row_idx': [THE_ROW_INDEX], 'column_idx': [THE_COLUMN_INDEX], 'text': [THE_TEXT_TO_WRITE]}"
)

def construct_action(work_dir, args: dict, py_file_path='/apps/excel_app/excel_set_cell.py'):
    text = str(args["text"])
    if text == 'None':
        text = '[None]'
    return build_command(py_file_path, file_path=args["file_path"], text=text,
                         row_idx=args["row_idx"], column_idx=args["column_idx"],
                         sheet_name=args.get("sheet_name"))


def excel_set_cell(file_path, text, row_idx, column_idx, sheet_name=None):
    if text is True:
        text = ''
    try:
        if os.path.exists(file_path):
            workbook = openpyxl.load_workbook(file_path)
        else:
            workbook = openpyxl.Workbook()
        if sheet_name is None:
            sheet = workbook.active
        else:
            try:
                sheet = workbook[sheet_name]
            except KeyError:
                sheet = workbook.create_sheet(title=sheet_name)
        sheet.cell(row=int(row_idx), column=int(column_idx), value=text)
        workbook.save(file_path)
        return True
    except Exception as e:
        print(e)
        return False


def main(file_path, text, row_idx, column_idx, sheet_name=None):
    if not os.path.exists(file_path):
        return f"OBSERVATION: The file {file_path} does not exist. Failed to write to the file."
    success = excel_set_cell(file_path, text, row_idx, column_idx, sheet_name)
    if success:
        return f'OBSERVATION: Successfully write text to {file_path}'
    return f'OBSERVATION: Failed to write text to {file_path}'


if __name__ == '__main__':
    fire.Fire(main)
