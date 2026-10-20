import os
import shutil
import sys
import tempfile

import openpyxl
import pytest
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import apps  
import run_benchmark as rb  
from apps.excel_app.excel_delete_cell import excel_delete_cell  
from apps.email_app import email_send_email, email_list_emails  # noqa: E402
from apps.word_app.word_read_file import word_read_file_into_string  # noqa: E402
from core.env import WorldBenchEnv  # noqa: E402
from core.sandbox import LocalSandbox  # noqa: E402
from evaluation import EVALUATION_FUNCTIONS as E  # noqa: E402


@pytest.fixture
def testbed():
    root = tempfile.mkdtemp()
    os.makedirs(os.path.join(root, "data"))
    os.makedirs(os.path.join(root, "emails"))
    os.makedirs(os.path.join(root, "calendar"))
    os.environ["TESTBED_ROOT"] = root
    yield root
    shutil.rmtree(root, ignore_errors=True)


def _workbook(path, cells):
    wb = openpyxl.Workbook()
    for (row, col), value in cells.items():
        wb.active.cell(row, col, value)
    wb.save(path)


def test_delete_cell_clears_the_cell(testbed):
    path = os.path.join(testbed, "data", "a.xlsx")
    _workbook(path, {(1, 1): "KEEPME"})
    assert excel_delete_cell(path, 1, 1)
    assert openpyxl.load_workbook(path).active.cell(1, 1).value is None


def test_not_contain_rejects_any_forbidden_keyword(testbed):
    path = os.path.join(testbed, "data", "n.txt")
    open(path, "w", encoding="utf-8").write("this document is private")
    args = {"doc_type": "txt", "file": "./data/n.txt", "keywords": ["private", "confidential"]}
    assert E["evaluate_not_contain"](testbed, args) is False
    args["keywords"] = ["secret", "confidential"]
    assert E["evaluate_not_contain"](testbed, args) is True


def test_excel_cell_value_is_exact(testbed):
    path = os.path.join(testbed, "data", "s.xlsx")
    _workbook(path, {(1, 1): "foobar", (2, 1): 50, (3, 1): 5})
    def check(row, value):
        return E["evaluate_excel_cell_value"](
            testbed, {"file": "./data/s.xlsx", "matches": [{"row": row, "col": 1, "value": value}]}
        )
    assert check(1, "foo") is False
    assert check(1, "foobar") is True
    assert check(2, "5") is False
    assert check(3, "5") is True
    assert check(3, 5) is True


def test_word_reader_sees_tables_and_headers(testbed):
    path = os.path.join(testbed, "data", "t.docx")
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE ONLY TEXT"
    doc.sections[0].header.paragraphs[0].text = "HEADER TEXT"
    doc.save(path)
    text = word_read_file_into_string(path)
    assert "TABLE ONLY TEXT" in text
    assert "HEADER TEXT" in text


def test_email_without_charset_is_readable(testbed):
    box = os.path.join(testbed, "emails", "bob")
    os.makedirs(box, exist_ok=True)
    open(os.path.join(box, "x.eml"), "w", encoding="utf-8").write(
        "From: a@x.com\nTo: bob@x.com\nSubject: S\nContent-Type: text/plain\n\nhello body\n"
    )
    out = email_list_emails.list_emails("bob", testbed, -1)
    assert "hello body" in out


def test_equal_subjects_are_both_kept(testbed):
    email_send_email.send_email("alice", "bob", "Report", "first")
    email_send_email.send_email("alice", "bob", "Report", "second")
    box = os.path.join(testbed, "emails", "bob")
    assert len(os.listdir(box)) == 2
    bodies = "".join(open(os.path.join(box, f), encoding="utf-8").read() for f in os.listdir(box))
    assert "first" in bodies and "second" in bodies


def test_subject_cannot_escape_the_mailbox(testbed):
    email_send_email.send_email("alice", "bob", "../../escaped", "x")
    assert not os.path.exists(os.path.join(testbed, "escaped.eml"))
    assert len(os.listdir(os.path.join(testbed, "emails", "bob"))) == 1


def test_quotes_survive_the_shell_round_trip(testbed):
    from apps.word_app import word_write_to_file

    sandbox = LocalSandbox(os.path.join(testbed, "sb"))
    sandbox.setup(source_testbed_dir=None, source_apps_dir="apps")
    target = os.path.join(sandbox.testbed_dir, "data", "quoted file.docx")
    Document().save(target)

    text = 'He said "yes" and it\'s fine'
    command = word_write_to_file.construct_action(
        "/testbed", {"file_path": "/testbed/data/quoted file.docx", "contents": text})
    code, out = sandbox.exec(command)
    assert code == 0, out
    assert text in word_read_file_into_string(target)


def test_shell_metacharacters_are_inert(testbed):
    from apps.email_app import email_send_email as sender
    command = sender.construct_action("/testbed", {
        "sender": "a", "recipient": "b", "subject": "S; touch /tmp/wb_pwned", "content": "x"})
    sandbox = LocalSandbox(os.path.join(testbed, "sb"))
    sandbox.setup(source_testbed_dir=None, source_apps_dir="apps")
    sandbox.exec(command)
    assert not os.path.exists("/tmp/wb_pwned")


def test_the_answer_never_lands_in_the_testbed(testbed):
    env = _env(testbed)
    answer = 'He said "yes" and it\'s fine'
    env.exec_action('{"app":"system","action":"finish_task","answer":%r}' % answer)
    assert env.answer == answer
    assert not os.path.exists(os.path.join(env.sandbox.testbed_dir, "data", "answer.txt"))
    output = os.path.join(testbed, "out")
    os.makedirs(output, exist_ok=True)
    env.dump_history(output)
    assert open(os.path.join(output, "answer.txt"), encoding="utf-8").read().strip() == answer


def test_credentials_are_not_reachable(testbed):
    os.environ["OPENAI_API_KEY"] = "sk-SECRET-PROBE"
    sandbox = LocalSandbox(os.path.join(testbed, "sb"))
    sandbox.setup(source_testbed_dir=None, source_apps_dir="apps",
                  openai_key="sk-SECRET-PROBE")
    _, out = sandbox.exec("echo $OPENAI_API_KEY; cat $SANDBOX_ROOT/openai_key.txt 2>/dev/null")
    assert "SECRET-PROBE" not in out


def _env(testbed):
    sandbox = LocalSandbox(os.path.join(testbed, "sb"))
    sandbox.setup(source_testbed_dir=None, source_apps_dir="apps")
    env = WorldBenchEnv(sandbox=sandbox, task="t", verbose=False)
    env.reset()
    return env


def test_info_is_reset_per_action(testbed):
    env = _env(testbed)
    env.exec_action("{'app':'shell','action':'command','command':'ls'}")
    assert env.info.get("action_executed") is True
    env.exec_action("{'app':'system','action':'switch_app','target_app':'excel'}")
    assert "action_executed" not in env.info


def test_unknown_app_is_rejected(testbed):
    env = _env(testbed)
    env.exec_action("{'app':'system','action':'switch_app','target_app':'does-not-exist'}")
    assert env.current_app is None
    assert "Unknown app" in env.observation
    env.exec_action("{'app':'system','action':'switch_app','target_app':'excel'}")
    assert env.get_available_actions()


def test_only_a_real_cd_changes_directory(testbed):
    env = _env(testbed)
    env.current_app = "shell"
    env.exec_action("{'app':'shell','action':'command','command':'cd data'}")
    assert env.sandbox.workdir.endswith("/data")
    before = env.sandbox.workdir
    env.exec_action("{'app':'shell','action':'command','command':'cdrom foo'}")
    assert env.sandbox.workdir == before


def test_not_contain_fails_on_a_missing_target(testbed):
    args = {"doc_type": "txt", "file": "./data/absent.txt", "keywords": ["secret"]}
    assert E["evaluate_not_contain"](testbed, args) is False
    assert E["evaluate_contain"](testbed, args) is False


def test_contain_still_normalises_thousands_separators(testbed):
    open(os.path.join(testbed, "data", "m.txt"), "w", encoding="utf-8").write("total 1,250")
    args = {"doc_type": "txt", "file": "./data/m.txt", "keywords": ["1250"]}
    assert E["evaluate_contain"](testbed, args) is True


def test_history_truncation_keeps_the_action(testbed):
    from core.policy import LLMPolicy

    policy = LLMPolicy.__new__(LLMPolicy)
    policy.max_history_chars = 300

    class Env:
        history = [({"app": "pdf", "action": "read_file"}, "X" * 5000)]

    rendered = policy._format_history(Env())
    assert "read_file" in rendered
    assert "observation truncated" in rendered
    assert len(rendered) <= 300


def test_app_action_requires_an_active_app(testbed):
    env = _env(testbed)
    env.exec_action("{'app':'excel','action':'create_new_file','file_path':'/testbed/data/x.xlsx'}")
    assert "not active" in env.observation
    assert not os.path.exists(os.path.join(env.sandbox.testbed_dir, "data", "x.xlsx"))
    env.exec_action("{'app':'system','action':'switch_app','target_app':'excel'}")
    env.exec_action("{'app':'excel','action':'create_new_file','file_path':'/testbed/data/x.xlsx'}")
    assert os.path.exists(os.path.join(env.sandbox.testbed_dir, "data", "x.xlsx"))


def test_mount_prefixes_respect_path_boundaries(testbed):
    sandbox = LocalSandbox(os.path.join(testbed, "sb"))
    sandbox.setup(source_testbed_dir=None, source_apps_dir="apps")
    assert sandbox._to_real_path("/testbed/a.txt").startswith(sandbox.testbed_dir)
    assert sandbox._to_real_path("/testbed2/a.txt") == "/testbed2/a.txt"
    assert sandbox._to_real_path("/apps2/x.py") == "/apps2/x.py"
    assert sandbox._rewrite_command("cat /testbed2/a.txt") == "cat /testbed2/a.txt"
    assert sandbox.testbed_dir in sandbox._rewrite_command("cat /testbed/a.txt")


def test_email_headers_preserve_the_supplied_domain(testbed):
    email_send_email.send_email("alice@university.ac.uk", "bob@nhs.uk", "Ref", "body")
    box = os.path.join(testbed, "emails", "bob")
    headers = open(os.path.join(box, os.listdir(box)[0]), encoding="utf-8").read()
    assert "alice@university.ac.uk" in headers
    assert "bob@nhs.uk" in headers
    assert os.path.isdir(os.path.join(testbed, "emails", "alice"))


def test_cjk_matching_ignores_missing_spacing_and_line_breaks(testbed):
    from evaluation.evaluators import _evaluate_contain_text

    extracted = "重庆立交\n桥墩P4\n边坡位移"
    args = {"doc_type": "pdf", "keywords": ["桥墩 P4", "立交桥墩"]}
    assert _evaluate_contain_text(extracted, args) is True
    assert _evaluate_contain_text("unrelated", args) is False


def test_evaluators_return_false_on_unreadable_files(testbed):
    data = os.path.join(testbed, "data")
    open(os.path.join(data, "corrupt.xlsx"), "wb").write(b"not a workbook")
    open(os.path.join(data, "corrupt.docx"), "wb").write(b"PK\x03\x04garbage")
    open(os.path.join(testbed, "calendar", "bad.ics"), "w").write("NOT AN ICS")

    assert E["evaluate_contain"](testbed, {"doc_type": "xlsx", "file": "./data/corrupt.xlsx",
                                           "keywords": ["x"]}) is False
    assert E["evaluate_not_contain"](testbed, {"doc_type": "docx", "file": "./data/corrupt.docx",
                                               "keywords": ["x"]}) is False
    assert E["evaluate_excel_cell_value"](testbed, {"file": "./data/corrupt.xlsx",
                                                    "matches": [{"row": 1, "col": 1, "value": "x"}]}) is False
    assert E["evaluate_calendar_no_overlap"](testbed, {"username": "bad"}) is False
    assert E["evaluate_calendar_no_overlap"](testbed, {"username": "ghost"}) is False


def test_file_exist_requires_a_file_not_a_directory(testbed):
    assert E["evaluate_file_exist"](testbed, {"file": "./data"}) is False
    open(os.path.join(testbed, "data", "out.txt"), "w").write("x")
    assert E["evaluate_file_exist"](testbed, {"file": "./data/out.txt"}) is True


def test_missing_sheet_fails_cleanly(testbed):
    path = os.path.join(testbed, "data", "s.xlsx")
    _workbook(path, {(1, 1): "x"})
    assert E["evaluate_excel_cell_value"](testbed, {"file": "./data/s.xlsx", "sheet_name": "Nope",
                                                    "matches": [{"row": 1, "col": 1, "value": "x"}]}) is False


def test_schema_errors_still_count_as_malformed(testbed):
    env = _env(testbed)
    env.exec_action("{'app':'system','action':'switch_app','target_app':'excel'}")
    for action in ("prose with no json",
                   '{"app":"excel","action":"set_cell","file_path":"/testbed/a.xlsx","row":1,"col":1,"text":"x"}',
                   '{"app":"excel","action":"write_cell","file_path":"/testbed/a.xlsx"}'):
        before = env.n_malformed
        env.exec_action(action)
        assert env.n_malformed == before + 1, action
    before = env.n_protocol_violations
    env.exec_action('{"app":"word","action":"read_file","file_path":"/testbed/a.docx"}')
    assert env.n_protocol_violations == before + 1
    assert env.n_malformed == 3


def test_calendar_event_carries_a_description(testbed):
    from apps.calendar_app.calendar_create_event import create_event

    assert create_event("Anna", "Sopralluogo", "2026-04-07 09:00:00", "2026-04-07 11:00:00",
                        description="Finestra ZTL 06:30-09:30")
    ics = open(os.path.join(testbed, "calendar", "Anna.ics"), encoding="utf-8").read()
    assert "06:30-09:30" in ics
    assert "test event" not in ics


def test_llm_app_is_off_by_default():
    assert "llm" not in apps.AVAILABLE_APPS
    assert "llm" in apps.ALL_APPS


def test_preservation_matches_targets_exactly():
    source = tempfile.mkdtemp()
    os.makedirs(os.path.join(source, "data"))
    open(os.path.join(source, "foo.txt"), "w").write("root level")
    open(os.path.join(source, "data", "foo.txt"), "w").write("target")
    result_root = tempfile.mkdtemp()
    result = os.path.join(result_root, "tb")
    shutil.copytree(source, result)
    config = {"evaluation": [{"function": "evaluate_contain",
                              "args": {"file": "./data/foo.txt"}}]}

    report = rb.compute_preservation(source, result, config)
    assert report["preserved"] is True and report["n_non_target"] == 1

    open(os.path.join(result, "foo.txt"), "w").write("TAMPERED")
    report = rb.compute_preservation(source, result, config)
    assert report["preserved"] is False and report["n_modified"] == 1

    shutil.rmtree(source, ignore_errors=True)
    shutil.rmtree(result_root, ignore_errors=True)


def _preservation_scenario(mutate):
    """Build a testbed with one target and two bystanders, then apply mutate."""
    source = tempfile.mkdtemp()
    os.makedirs(os.path.join(source, "data"))
    for name in ("target.xlsx", "other.txt", "notes.txt"):
        open(os.path.join(source, "data", name), "w").write("original")
    result_root = tempfile.mkdtemp()
    result = os.path.join(result_root, "tb")
    shutil.copytree(source, result)
    mutate(result)
    config = {"evaluation": [{"function": "evaluate_excel_cell_value",
                              "args": {"file": "./data/target.xlsx", "matches": []}}]}
    report = rb.compute_preservation(source, result, config)
    shutil.rmtree(source, ignore_errors=True)
    shutil.rmtree(result_root, ignore_errors=True)
    return report


def test_modifying_a_bystander_breaks_preservation():
    report = _preservation_scenario(
        lambda tb: open(os.path.join(tb, "data", "other.txt"), "w").write("edited"))
    assert report["preserved"] is False
    assert report["n_modified"] == 1 and report["acted"] is True


def test_deleting_a_bystander_breaks_preservation():
    report = _preservation_scenario(
        lambda tb: os.remove(os.path.join(tb, "data", "notes.txt")))
    assert report["preserved"] is False
    assert report["n_deleted"] == 1 and report["acted"] is True


def test_editing_the_target_leaves_preservation_intact():
    report = _preservation_scenario(
        lambda tb: open(os.path.join(tb, "data", "target.xlsx"), "w").write("edited"))
    assert report["preserved"] is True
    assert report["n_modified"] == 0 and report["acted"] is True


def test_a_new_artefact_is_counted_but_does_not_break_preservation():
    report = _preservation_scenario(
        lambda tb: open(os.path.join(tb, "data", "scratch.docx"), "w").write("new"))
    assert report["preserved"] is True
    assert report["n_created"] == 1 and report["acted"] is True


def test_an_agent_written_answer_file_is_an_artefact_like_any_other():
    report = _preservation_scenario(
        lambda tb: open(os.path.join(tb, "data", "answer.txt"), "w").write("None\n"))
    assert report["preserved"] is True
    assert report["n_created"] == 1 and report["acted"] is True


def test_renaming_a_bystander_breaks_preservation():
    def rename(tb):
        os.rename(os.path.join(tb, "data", "notes.txt"),
                  os.path.join(tb, "data", "notes_old.txt"))
    report = _preservation_scenario(rename)
    assert report["preserved"] is False
    assert report["n_deleted"] == 1 and report["n_created"] == 1


def test_emptying_a_bystander_breaks_preservation():
    report = _preservation_scenario(
        lambda tb: open(os.path.join(tb, "data", "other.txt"), "w").close())
    assert report["preserved"] is False and report["n_modified"] == 1


def test_deleting_the_target_leaves_preservation_intact():
    report = _preservation_scenario(
        lambda tb: os.remove(os.path.join(tb, "data", "target.xlsx")))
    assert report["preserved"] is True and report["acted"] is True


def test_a_bystander_swapped_for_a_symlink_is_not_scorable():
    def swap(tb):
        victim = os.path.join(tb, "data", "other.txt")
        decoy = os.path.join(tb, "data", ".decoy")
        shutil.copyfile(victim, decoy)
        os.remove(victim)
        os.symlink(decoy, victim)
    with pytest.raises(rb.PreservationUnscorable):
        _preservation_scenario(swap)


def test_every_declared_target_is_exempt():
    config = {"evaluation": [
        {"function": "evaluate_contain", "args": {"file": "./data/a.txt"}},
        {"function": "evaluate_contain", "args": {"file": "/testbed/data/b.txt"}},
        {"function": "evaluate_calendar_no_overlap", "args": {"username": "Anna"}},
    ]}
    assert rb._eval_target_files(config) == {
        os.path.join("data", "a.txt"),
        os.path.join("data", "b.txt"),
        os.path.join("calendar", "Anna.ics"),
        os.path.join("data", "Anna.ics"),
    }


def test_a_missing_pristine_state_is_never_a_pass():
    with pytest.raises(rb.PreservationUnscorable):
        rb.compute_preservation("/nonexistent/source", tempfile.mkdtemp(), {"evaluation": []})


def test_a_missing_final_state_is_never_a_pass():
    with pytest.raises(rb.PreservationUnscorable):
        rb.compute_preservation(tempfile.mkdtemp(), "/nonexistent/result", {"evaluation": []})


def test_listing_an_absent_calendar_creates_nothing(testbed):
    from apps.calendar_app.calendar_list_events import list_events

    before = rb._list_files_rel(testbed)
    assert list_events("Domenico") == ""
    assert not os.path.exists(os.path.join(testbed, "calendar", "Domenico.ics"))
    assert rb._list_files_rel(testbed) == before


def test_reading_an_absent_mailbox_creates_nothing(testbed):
    from apps.email_app.email_read_email import read_email

    before = rb._list_files_rel(testbed)
    assert list(email_list_emails.list_emails("nobody")) == list("No emails found.")
    read_email("nobody", "missing.eml")
    assert not os.path.isdir(os.path.join(testbed, "emails", "nobody"))
    assert rb._list_files_rel(testbed) == before


def test_an_inert_agent_preserves_vacuously():
    report = _preservation_scenario(lambda tb: None)
    assert report["preserved"] is True and report["acted"] is False


def test_prose_quotes_do_not_hide_the_action():
    from core.env import parse_action

    payload = '{"app":"excel","action":"set_cell","file_path":"/x.xlsx",' \
              '"row_idx":2,"column_idx":2,"text":"No"}'
    for preamble in ("Aggiorno lo stato dell'ordine.\n",
                     "I'll update the cell now.\n",
                     "Je mets a jour l'indemnite.\n",
                     ""):
        assert parse_action(preamble + payload)["action"] == "set_cell", preamble


def test_calendar_keywords_survive_folding_and_escaping(testbed):
    from apps.calendar_app.calendar_create_event import create_event

    assert create_event("cal", "Buffer", "2026-03-19 08:00:00", "2026-03-19 12:00:00",
                        description="x" * 60 + " window 08:00, 12:00; ring road")
    assert E["evaluate_contain"](testbed, {"doc_type": "ics",
                                           "file": "calendar/cal.ics",
                                           "keywords": ["08:00, 12:00", "ring road"]})


def test_calendar_target_is_derived_from_the_username():
    config = {"evaluation": [{"function": "evaluate_calendar_no_overlap",
                              "args": {"username": "Giulia"}}]}
    assert rb._eval_target_files(config) == {
        os.path.join("calendar", "Giulia.ics"),
        os.path.join("data", "Giulia.ics"),
    }


def test_observation_budget(testbed):
    sandbox = LocalSandbox(os.path.join(testbed, "sb"))
    sandbox.setup(source_testbed_dir=None, source_apps_dir="apps")
    env = WorldBenchEnv(sandbox=sandbox, task="t", verbose=False, max_obs_chars=50)
    env.reset()
    env.current_app = "shell"
    observation, _, _, _ = env.step(
        "{'app':'shell','action':'command','command':'head -c 5000 /dev/zero | tr \\\\0 a'}"
    )
    assert len(str(observation)) < 200
    assert "truncated" in str(observation)
    assert env.history[-1][1] == observation


def test_target_paths_normalise_to_one_form():
    forms = ["./data/x.txt", "data/x.txt", "/testbed/data/x.txt",
             "testbed/data/x.txt", "\\testbed\\data\\x.txt"]
    assert {rb._normalise_target(f) for f in forms} == {os.path.join("data", "x.txt")}


def test_a_missing_pristine_testbed_is_never_preserved():
    result = tempfile.mkdtemp()
    with pytest.raises(rb.PreservationUnscorable):
        rb.compute_preservation(os.path.join(result, "absent"), result, {"evaluation": []})
    with pytest.raises(rb.PreservationUnscorable):
        rb.compute_preservation(None, result, {"evaluation": []})
    shutil.rmtree(result, ignore_errors=True)


def test_a_missing_final_testbed_is_never_preserved():
    source = tempfile.mkdtemp()
    with pytest.raises(rb.PreservationUnscorable):
        rb.compute_preservation(source, os.path.join(source, "absent"), {"evaluation": []})
    shutil.rmtree(source, ignore_errors=True)


def test_an_unreadable_file_does_not_become_an_empty_one():
    source = tempfile.mkdtemp()
    os.makedirs(os.path.join(source, "data"))
    path = os.path.join(source, "data", "locked.txt")
    open(path, "w").write("secret")
    os.chmod(path, 0o000)
    try:
        if os.access(path, os.R_OK):
            pytest.skip("running as a user that bypasses file permissions")
        with pytest.raises(PermissionError):
            rb._list_files_rel(source)
    finally:
        os.chmod(path, 0o644)
        shutil.rmtree(source, ignore_errors=True)


def test_a_symlinked_bystander_is_not_silently_equivalent():
    source = tempfile.mkdtemp()
    os.makedirs(os.path.join(source, "data"))
    open(os.path.join(source, "data", "other.txt"), "w").write("original")
    result_root = tempfile.mkdtemp()
    result = os.path.join(result_root, "tb")
    shutil.copytree(source, result)
    decoy = os.path.join(result_root, "decoy.txt")
    open(decoy, "w").write("original")
    target = os.path.join(result, "data", "other.txt")
    os.remove(target)
    os.symlink(decoy, target)
    with pytest.raises(rb.PreservationUnscorable):
        rb.compute_preservation(source, result, {"evaluation": []})
    shutil.rmtree(source, ignore_errors=True)
    shutil.rmtree(result_root, ignore_errors=True)




def _write(path, text):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


@pytest.mark.parametrize("text, keyword, expected", [
    ("Please  ARRIVE thirty, minutes early!", "arrive thirty minutes early", True),
    ("arrive – thirty minutes early", "arrive thirty minutes early", True),
    ("arrive early thirty minutes", "arrive thirty minutes early", False),
    ("alpha " + "z" * 30 + " beta", "alpha beta", True),
    ("alpha " + "z" * 50 + " beta", "alpha beta", False),
    ("alpha\nbeta", "alpha beta", False),
    ("registered", "register", False),
    ("Total 1,250 units", "1250", True),
    ("Total 12500 units", "1250", False),
    ("Invoice INV-1250", "1250", False),
    ("cost 12.50", "12.5", False),
    ("2026-03-19", "2026 03 19", False),
    ("it’s fine", "it's fine", True),
    ("50 %", "50%", True),
    ("50", "50%", False),
])
def test_flexible_matching_rules(text, keyword, expected):
    from evaluation import matching

    assert matching.contains(text, keyword) is expected


def test_empty_verification_is_rejected(testbed):
    _write(os.path.join(testbed, "data", "a.txt"), "text")
    with pytest.raises(ValueError):
        E["evaluate_contain"](testbed, {"doc_type": "txt", "file": "./data/a.txt", "keywords": []})
    with pytest.raises(ValueError):
        E["evaluate_not_contain"](testbed, {"doc_type": "txt", "file": "./data/a.txt"})
    _workbook(os.path.join(testbed, "data", "s.xlsx"), {(1, 1): "x"})
    with pytest.raises(ValueError):
        E["evaluate_excel_cell_value"](testbed, {"file": "./data/s.xlsx", "matches": []})


def test_excel_coordinates_do_not_satisfy_numeric_keywords(testbed):
    _workbook(os.path.join(testbed, "data", "s.xlsx"), {(1, 1): "alpha", (2, 2): 40})
    args = {"doc_type": "xlsx", "file": "./data/s.xlsx", "keywords": ["2"]}
    assert E["evaluate_contain"](testbed, args) is False
    args["keywords"] = ["40"]
    assert E["evaluate_contain"](testbed, args) is True


def _events(testbed, name, events, folder="calendar"):
    from apps.calendar_app.calendar_create_event import create_event

    for event in events:
        assert create_event(name, *event[:3], description=event[3], location=event[4])
    return os.path.join(testbed, folder, f"{name}.ics")


def test_calendar_event_requires_every_property_on_one_event(testbed):
    _events(testbed, "ann", [
        ("Team Sync", "2026-03-19 09:00:00", "2026-03-19 10:00:00", "agenda: budget", "Room 4"),
        ("Site Visit", "2026-03-20 14:00:00", "2026-03-20 16:00:00", "bring helmet", "Depot 7"),
    ])

    def check(**props):
        return E["evaluate_calendar_event"](testbed, {"username": "ann", **props})

    assert check(summary_keywords=["team  sync!"], date="2026-03-19",
                 time_start="2026-03-19 09:00:00", time_end="2026-03-19 10:00:00",
                 description_keywords=["budget agenda"]) is False
    assert check(summary_keywords=["team sync"], date="2026-03-19",
                 time_start="2026-03-19 09:00:00", time_end="2026-03-19 10:00:00",
                 description_keywords=["budget"], location_keywords=["room 4"]) is True
    assert check(summary_keywords=["Team Sync"], description_keywords=["helmet"]) is False
    assert check(summary_keywords=["Site Visit"], date="2026-03-19") is False
    assert check(summary_keywords=["Site Visit"], time_start="2026-03-20 14:01:00") is False
    assert check(summary_keywords=["Site Visit"], location_keywords=["Depot 70"]) is False
    with pytest.raises(ValueError):
        check()


def test_calendar_event_compares_aware_times_consistently(testbed):
    _write(os.path.join(testbed, "calendar", "bo.ics"), (
        "BEGIN:VCALENDAR\r\nVERSION:2.0\r\nPRODID:-//t//\r\nBEGIN:VEVENT\r\nUID:1\r\n"
        "DTSTAMP:20260301T000000Z\r\nSUMMARY:Call\r\nDTSTART:20260319T080000Z\r\n"
        "DTEND:20260319T090000Z\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n"))
    args = {"username": "bo", "summary_keywords": ["Call"], "time_start": "2026-03-19T09:00:00+01:00",
            "time_end": "2026-03-19T10:00:00+01:00"}
    assert E["evaluate_calendar_event"](testbed, args) is True
    assert E["evaluate_calendar_event"](testbed, {**args, "timezone": "Europe/Rome",
                                                  "time_start": "09:00", "date": "2026-03-19"}) is True
    assert E["evaluate_calendar_event"](testbed, {**args, "time_start": "2026-03-19T09:00:00+00:00"}) is False


def test_calendar_app_edits_an_existing_calendar_in_data(testbed):
    from apps.calendar_app.calendar_create_event import create_event
    from apps.calendar_app.calendar_list_events import list_events

    assert create_event("planner", "Old", "2026-03-01 08:00:00", "2026-03-01 09:00:00")
    os.replace(os.path.join(testbed, "calendar", "planner.ics"), os.path.join(testbed, "data", "planner.ics"))
    assert create_event("planner", "New", "2026-03-02 08:00:00", "2026-03-02 09:00:00")
    assert not os.path.exists(os.path.join(testbed, "calendar", "planner.ics"))
    assert "Old" in list_events("planner.ics") and "New" in list_events("planner")
    args = {"file": "./data/planner.ics", "summary_keywords": ["new"], "date": "2026-03-02"}
    assert E["evaluate_calendar_event"](testbed, args) is True


def test_calendar_files_resolve_from_calendar_or_data(testbed):
    for folder in ("calendar", "data"):
        _write(os.path.join(testbed, folder, f"{folder}_user.ics"), (
            "BEGIN:VCALENDAR\r\nVERSION:2.0\r\nPRODID:-//t//\r\nBEGIN:VEVENT\r\nUID:1\r\n"
            "DTSTAMP:20260301T000000Z\r\nSUMMARY:A\r\nDTSTART:20260319T080000Z\r\n"
            "DTEND:20260319T100000Z\r\nEND:VEVENT\r\nBEGIN:VEVENT\r\nUID:2\r\n"
            "DTSTAMP:20260301T000000Z\r\nSUMMARY:B\r\nDTSTART:20260319T090000Z\r\n"
            "DTEND:20260319T110000Z\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n"))
        assert E["evaluate_calendar_no_overlap"](testbed, {"username": f"{folder}_user"}) is False


def _send(testbed, sender, recipient, subject, body):
    os.environ["TESTBED_ROOT"] = testbed
    email_send_email.send_email(sender, recipient, subject, body)


def test_sent_email_ignores_preexisting_and_undelivered_messages(testbed):
    _send(testbed, "ann", "bob@nhs.uk", "Old report", "attendance figures attached")
    initial = tempfile.mkdtemp()
    shutil.copytree(testbed, os.path.join(initial, "tb"))
    initial = os.path.join(initial, "tb")
    args = {"recipient": "bob@nhs.uk", "subject_keywords": ["report"],
            "keywords": ["attendance figures"]}

    assert E["evaluate_sent_email"](testbed, args, initial_dir=initial) is False
    assert E["evaluate_sent_email"](testbed, args) is True

    _write(os.path.join(testbed, "data", "draft.txt"),
           "To: bob@nhs.uk\nSubject: New report\n\nattendance figures")
    assert E["evaluate_sent_email"](testbed, args, initial_dir=initial) is False

    _send(testbed, "ann", "carol@nhs.uk", "New report", "attendance figures")
    assert E["evaluate_sent_email"](testbed, args, initial_dir=initial) is False

    _send(testbed, "ann", "bob@nhs.uk", "New Report!", "The attendance  figures are ready")
    assert E["evaluate_sent_email"](testbed, args, initial_dir=initial) is True
    assert E["evaluate_sent_email"](testbed, {**args, "recipient": "bob@other.org"},
                                    initial_dir=initial) is False
    assert E["evaluate_sent_email"](testbed, {**args, "keywords": ["attendance", "budget"]},
                                    initial_dir=initial) is False
    with pytest.raises(ValueError):
        E["evaluate_sent_email"](testbed, {"recipient": "bob@nhs.uk"}, initial_dir=initial)
    shutil.rmtree(os.path.dirname(initial), ignore_errors=True)


def test_judge_verdict_must_be_an_unambiguous_yes_or_no():
    from evaluation.evaluators import _parse_verdict

    assert _parse_verdict("YES") is True
    assert _parse_verdict(" no. ") is False
    assert _parse_verdict("**Yes**") is True
    for answer in ("Yes, but the greeting is missing", "No, not yes", "maybe", "", None):
        with pytest.raises(ValueError):
            _parse_verdict(answer)


def test_empty_evaluation_never_passes():
    passed, details = rb.evaluate_task({"evaluation": []}, tempfile.mkdtemp())
    assert passed is False and "error" in details[0]
    passed, details = rb.evaluate_task({}, tempfile.mkdtemp())
    assert passed is False


def test_evaluator_errors_are_reported_apart_from_failures(testbed):
    _write(os.path.join(testbed, "data", "a.txt"), "hello")
    config = {"evaluation": [
        {"function": "evaluate_contain",
         "args": {"doc_type": "txt", "file": "./data/a.txt", "keywords": ["absent"]}},
        {"function": "evaluate_contain",
         "args": {"doc_type": "txt", "file": "./data/a.txt", "keywords": []}},
        {"function": "evaluate_nothing", "args": {}},
    ]}
    passed, details = rb.evaluate_task(config, testbed)
    assert passed is False
    assert ["error" in d for d in details] == [False, True, True]
    assert all(d["passed"] is False for d in details)


def test_glob_targets_exempt_matching_files_only():
    source = tempfile.mkdtemp()
    _write(os.path.join(source, "data", "keep.txt"), "original")
    result = os.path.join(tempfile.mkdtemp(), "tb")
    shutil.copytree(source, result)
    _write(os.path.join(result, "data", "email_to_ops_1.eml"), "new")
    config = {"target_files": [{"path": "data/email_to_ops_*.eml", "change": "create"}]}
    assert rb.compute_preservation(source, result, config)["preserved"] is True
    assert "data/email_to_ops_*.eml" in rb._eval_target_files(config)


def test_preservation_targets_combine_every_declared_source():
    config = {
        "target_files": ["data/a.docx", {"path": "./data/b.xlsx"}],
        "output_files": ["report.txt", {"path": "/testbed/data/c.txt"}],
        "evaluation": [{"function": "evaluate_contain", "args": {"file": "./data/d.txt"}}],
    }
    assert rb._eval_target_files(config) == {
        os.path.join("data", name) for name in ("a.docx", "b.xlsx", "c.txt", "d.txt", "report.txt")
    } | {"report.txt"}
    assert rb._eval_target_files({"evaluation": []}) == set()


def test_declared_output_files_are_exempt_from_preservation():
    source = tempfile.mkdtemp()
    _write(os.path.join(source, "data", "notes.txt"), "original")
    _write(os.path.join(source, "data", "bystander.txt"), "original")
    result = os.path.join(tempfile.mkdtemp(), "tb")
    shutil.copytree(source, result)
    _write(os.path.join(result, "data", "notes.txt"), "edited")
    config = {"output_files": ["notes.txt"], "evaluation": []}
    assert rb.compute_preservation(source, result, config)["preserved"] is True
    _write(os.path.join(result, "data", "bystander.txt"), "edited")
    assert rb.compute_preservation(source, result, config)["preserved"] is False
